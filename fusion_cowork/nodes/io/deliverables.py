"""Deliverables IO 节点 — Excel/PPT 交付物生成 (issue #151 #152)。

xlsx_write: openpyxl 多 sheet 工作簿, "=..." 字符串按公式写入, 表头样式 +
条件格式。pptx_write: python-pptx 从 JSON spec 生成可编辑 deck (标题/内容/
图表/表格), period_compare 自动生成环比变动页 (>threshold% 打标)。
依赖 openpyxl / python-pptx (extra [office]), 缺失时显式失败不静默。
"""

from __future__ import annotations

import json
import logging
import time
from pathlib import Path
from typing import Any, Dict, List

from ...engine.node import (
    BaseNode,
    NodeCategory,
    NodeResult,
    NodeStatus,
    register_node,
)

logger = logging.getLogger(__name__)

_HEADER_FILL_HEX = "FF305496"
_FLAG_FILL_HEX = "FFFFC7CE"


def _resolve_output_path(scope_mgr, output_path: str, file_name: str, ext: str, overwrite: bool):
    out_dir = Path(output_path).expanduser()
    if not scope_mgr.ensure_allowed(out_dir):
        return None, f"沙箱拒绝越界写入: {out_dir}"
    out_dir.mkdir(parents=True, exist_ok=True)
    date_str = time.strftime("%Y%m%d_%H%M%S")
    actual_name = file_name.replace("{date}", date_str)
    out_path = out_dir / f"{actual_name}{ext}"
    if out_path.exists() and not overwrite:
        out_path = out_dir / f"{actual_name}_{int(time.time())}{ext}"
    if not scope_mgr.ensure_allowed(out_path):
        return None, f"沙箱拒绝越界写入: {out_path}"
    return out_path, ""


def _parse_spec(node: BaseNode, inputs: Dict[str, Any], spec_key: str = "spec") -> Any:
    raw = inputs.get(spec_key, node.config.params.get(spec_key))
    if raw is None:
        return None
    if isinstance(raw, (dict, list)):
        return raw
    try:
        return json.loads(raw)
    except (json.JSONDecodeError, TypeError) as e:
        logger.warning(f"spec JSON 解析失败: {e}")
        return None


@register_node
class XlsxWriteNode(BaseNode):
    """xlsx_write — 生成多 sheet Excel 工作簿 (公式/条件格式/表头样式)。"""

    name = "xlsx_write"
    display_name = "写入Excel"
    category = NodeCategory.IO
    description = "从 JSON spec 生成 Excel 工作簿: 多 sheet、公式、条件格式、表头样式"
    icon = "📊"
    default_label = "写入Excel"

    inputs = [
        {"key": "spec", "label": "工作簿 Spec", "type": "object", "optional": True},
        {"key": "output_path", "label": "输出目录", "type": "string", "optional": True},
    ]
    outputs = [
        {"key": "file_path", "label": "输出文件路径", "type": "string"},
        {"key": "sheet_count", "label": "Sheet 数", "type": "integer"},
        {"key": "row_count", "label": "数据行数", "type": "integer"},
    ]

    def get_params_schema(self) -> Dict[str, Any]:
        return {
            "type": "object",
            "properties": {
                "output_path": {"type": "string", "default": "~/Desktop/Output", "description": "输出目录"},
                "file_name": {"type": "string", "default": "workbook_{date}", "description": "文件名（不含扩展名）"},
                "overwrite": {"type": "boolean", "default": False},
                "spec": {"type": "string", "description": "工作簿 spec JSON（也可由上游输入）"},
            },
        }

    async def execute(self, inputs: Dict[str, Any]) -> NodeResult:
        try:
            import openpyxl
            from openpyxl.formatting.rule import CellIsRule
            from openpyxl.styles import Alignment, Font, PatternFill
            from openpyxl.utils import get_column_letter
        except ImportError as e:
            return NodeResult(
                status=NodeStatus.FAILED,
                error=f"缺少依赖 openpyxl ({e}) — 安装: pip install -e '.[office]'",
                summary="xlsx_write 依赖缺失",
            )

        params = self.config.params
        spec = _parse_spec(self, inputs)
        if not isinstance(spec, dict) or not (spec.get("sheets") or []):
            return NodeResult(status=NodeStatus.FAILED, error="spec 缺失或无 sheets 定义")

        from ...security import get_scoped_folder_manager

        scope = get_scoped_folder_manager()
        out_path, err = _resolve_output_path(
            scope,
            inputs.get("output_path", params.get("output_path", "~/Desktop/Output")),
            params.get("file_name", "workbook_{date}"),
            ".xlsx",
            params.get("overwrite", False),
        )
        if out_path is None:
            return NodeResult(status=NodeStatus.FAILED, error=err, summary="授权文件夹外路径被拦截")

        _OP_MAP = {
            ">": "greaterThan",
            ">=": "greaterThanOrEqual",
            "<": "lessThan",
            "<=": "lessThanOrEqual",
            "==": "equal",
            "=": "equal",
            "!=": "notEqual",
        }

        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        total_rows = 0
        try:
            for sheet_spec in spec["sheets"]:
                name = str(sheet_spec.get("name") or f"Sheet{len(wb.sheetnames) + 1}")
                ws = wb.create_sheet(title=name[:31])
                headers = sheet_spec.get("headers") or []
                rows = sheet_spec.get("rows") or []
                if headers:
                    ws.append(headers)
                    if sheet_spec.get("header_style", True):
                        fill = PatternFill(start_color=_HEADER_FILL_HEX, end_color=_HEADER_FILL_HEX, fill_type="solid")
                        for col in range(1, len(headers) + 1):
                            cell = ws.cell(row=1, column=col)
                            cell.font = Font(bold=True, color="FFFFFFFF")
                            cell.fill = fill
                            cell.alignment = Alignment(horizontal="center")
                        ws.freeze_panes = "A2"
                for row in rows:
                    # "=..." 字符串由 openpyxl 原生按公式写入 (打开时重算)
                    ws.append(row)
                total_rows += len(rows)
                widths = sheet_spec.get("column_widths") or []
                for idx, width in enumerate(widths, start=1):
                    ws.column_dimensions[get_column_letter(idx)].width = width
                for rule_spec in sheet_spec.get("conditional") or []:
                    rng = rule_spec.get("range")
                    op = _OP_MAP.get(str(rule_spec.get("op", ">")), "greaterThan")
                    value = rule_spec.get("value", 0)
                    color = str(rule_spec.get("color", _FLAG_FILL_HEX))
                    if not rng:
                        continue
                    fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
                    ws.conditional_formatting.add(
                        rng,
                        CellIsRule(operator=op, formula=[str(value)], fill=fill),
                    )
            if spec.get("active_sheet") and spec["active_sheet"] in wb.sheetnames:
                wb.active = wb.sheetnames.index(spec["active_sheet"])
            wb.save(out_path)
        except Exception as e:
            logger.error(f"xlsx_write 写入失败: {e}")
            return NodeResult(status=NodeStatus.FAILED, error=f"工作簿写入失败: {e}", summary="写入失败")

        logger.info(f"xlsx_write 完成: {out_path} ({len(wb.sheetnames)} sheets, {total_rows} rows)")
        return NodeResult(
            status=NodeStatus.SUCCESS,
            data={
                "file_path": str(out_path),
                "sheet_count": len(wb.sheetnames),
                "row_count": total_rows,
                "size": out_path.stat().st_size,
            },
            summary=f"已生成 {out_path.name} ({len(wb.sheetnames)} sheets, {total_rows} 行)",
        )


_CHART_TYPES = {
    "column": "COLUMN_CLUSTERED",
    "bar": "BAR_CLUSTERED",
    "line": "LINE_MARKERS",
    "pie": "PIE",
}

_PPT_LAYOUT_TITLE = 0
_PPT_LAYOUT_CONTENT = 1
_PPT_LAYOUT_BLANK = 6


@register_node
class PptxWriteNode(BaseNode):
    """pptx_write — 从 JSON spec 生成 PowerPoint deck (标题/内容/图表/表格 + 环比变动页)。"""

    name = "pptx_write"
    display_name = "写入PPT"
    category = NodeCategory.IO
    description = "从 JSON spec 生成可编辑 PPT: 标题/内容/图表/表格版式, 支持环比变动自动打标"
    icon = "📽️"
    default_label = "写入PPT"

    inputs = [
        {"key": "spec", "label": "Deck Spec", "type": "object", "optional": True},
        {"key": "output_path", "label": "输出目录", "type": "string", "optional": True},
    ]
    outputs = [
        {"key": "file_path", "label": "输出文件路径", "type": "string"},
        {"key": "slide_count", "label": "页数", "type": "integer"},
    ]

    def get_params_schema(self) -> Dict[str, Any]:
        return {
            "type": "object",
            "properties": {
                "output_path": {"type": "string", "default": "~/Desktop/Output", "description": "输出目录"},
                "file_name": {"type": "string", "default": "deck_{date}", "description": "文件名（不含扩展名）"},
                "overwrite": {"type": "boolean", "default": False},
                "spec": {"type": "string", "description": "Deck spec JSON（也可由上游输入）"},
            },
        }

    def _period_compare_slides(self, cmp_spec: Dict[str, Any]) -> List[Dict[str, Any]]:
        metric = str(cmp_spec.get("metric", "指标"))
        categories = cmp_spec.get("categories") or []
        current = cmp_spec.get("current") or []
        previous = cmp_spec.get("previous") or []
        threshold = float(cmp_spec.get("threshold_pct", 10))
        title = str(cmp_spec.get("title", f"环比变动 (>{threshold:g}% 打标)"))
        headers = [metric, "上期", "本期", "变动%", "标记"]
        rows: List[List[Any]] = []
        for idx, cat in enumerate(categories):
            cur = current[idx] if idx < len(current) else None
            prev = previous[idx] if idx < len(previous) else None
            if not isinstance(cur, (int, float)) or not isinstance(prev, (int, float)) or prev == 0:
                rows.append([cat, prev, cur, "", ""])
                continue
            pct = (cur - prev) / abs(prev) * 100
            rows.append([cat, prev, cur, round(pct, 1), "⚠️ 超阈值" if abs(pct) > threshold else ""])
        return [
            {
                "layout": "content",
                "title": title,
                "bullets": [
                    f"{cat}: {prev} → {cur} ({row[3]}%) {row[4]}".rstrip() for cat, row in zip(categories, rows)
                ],
            },
            {"layout": "table", "title": title, "headers": headers, "rows": rows},
        ]

    async def execute(self, inputs: Dict[str, Any]) -> NodeResult:
        try:
            from pptx import Presentation
            from pptx.chart.data import CategoryChartData
            from pptx.enum.chart import XL_CHART_TYPE
            from pptx.util import Inches, Pt
        except ImportError as e:
            return NodeResult(
                status=NodeStatus.FAILED,
                error=f"缺少依赖 python-pptx ({e}) — 安装: pip install -e '.[office]'",
                summary="pptx_write 依赖缺失",
            )

        params = self.config.params
        spec = _parse_spec(self, inputs)
        if not isinstance(spec, dict):
            return NodeResult(status=NodeStatus.FAILED, error="spec 缺失或非法（需 JSON 对象）")

        slides: List[Dict[str, Any]] = list(spec.get("slides") or [])
        if isinstance(spec.get("period_compare"), dict):
            slides.extend(self._period_compare_slides(spec["period_compare"]))
        if not slides:
            return NodeResult(status=NodeStatus.FAILED, error="spec 无 slides 定义")
        deck_title = str(spec.get("title") or "")

        from ...security import get_scoped_folder_manager

        scope = get_scoped_folder_manager()
        out_path, err = _resolve_output_path(
            scope,
            inputs.get("output_path", params.get("output_path", "~/Desktop/Output")),
            params.get("file_name", "deck_{date}"),
            ".pptx",
            params.get("overwrite", False),
        )
        if out_path is None:
            return NodeResult(status=NodeStatus.FAILED, error=err, summary="授权文件夹外路径被拦截")

        prs = Presentation()
        try:
            if deck_title:
                slide = prs.slides.add_slide(prs.slide_layouts[_PPT_LAYOUT_TITLE])
                slide.shapes.title.text = deck_title
                if spec.get("subtitle"):
                    slide.placeholders[1].text = str(spec["subtitle"])
            for s_spec in slides:
                layout = str(s_spec.get("layout", "content"))
                if layout == "title":
                    slide = prs.slides.add_slide(prs.slide_layouts[_PPT_LAYOUT_TITLE])
                    slide.shapes.title.text = str(s_spec.get("title", ""))
                    if s_spec.get("subtitle"):
                        slide.placeholders[1].text = str(s_spec["subtitle"])
                elif layout == "chart" and isinstance(s_spec.get("chart"), dict):
                    slide = prs.slides.add_slide(prs.slide_layouts[_PPT_LAYOUT_BLANK])
                    if s_spec.get("title"):
                        box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.8))
                        box.text_frame.text = str(s_spec["title"])
                    chart_spec = s_spec["chart"]
                    chart_data = CategoryChartData()
                    chart_data.categories = [str(c) for c in chart_spec.get("categories") or []]
                    for series in chart_spec.get("series") or []:
                        chart_data.add_series(str(series.get("name", "series")), series.get("values") or [])
                    type_name = _CHART_TYPES.get(str(chart_spec.get("type", "column")), "COLUMN_CLUSTERED")
                    slide.shapes.add_chart(
                        getattr(XL_CHART_TYPE, type_name),
                        Inches(0.5),
                        Inches(1.2),
                        Inches(9),
                        Inches(5.5),
                        chart_data,
                    )
                elif layout == "table":
                    slide = prs.slides.add_slide(prs.slide_layouts[_PPT_LAYOUT_BLANK])
                    if s_spec.get("title"):
                        box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.8))
                        box.text_frame.text = str(s_spec["title"])
                    headers = s_spec.get("headers") or []
                    rows = s_spec.get("rows") or []
                    n_rows = len(rows) + (1 if headers else 0)
                    n_cols = max([len(headers)] + [len(r) for r in rows]) if (headers or rows) else 1
                    table = slide.shapes.add_table(
                        max(n_rows, 1), max(n_cols, 1), Inches(0.5), Inches(1.2), Inches(9), Inches(0.5)
                    ).table
                    if headers:
                        for c, h in enumerate(headers):
                            table.cell(0, c).text = str(h)
                    for r, row in enumerate(rows, start=1 if headers else 0):
                        for c, val in enumerate(row):
                            table.cell(r, c).text = str(val)
                else:
                    slide = prs.slides.add_slide(prs.slide_layouts[_PPT_LAYOUT_CONTENT])
                    slide.shapes.title.text = str(s_spec.get("title", ""))
                    body = slide.placeholders[1].text_frame
                    bullets = s_spec.get("bullets") or []
                    for i, bullet in enumerate(bullets):
                        para = body.paragraphs[0] if i == 0 else body.add_paragraph()
                        para.text = str(bullet)
                        para.font.size = Pt(18)
            prs.save(out_path)
        except Exception as e:
            logger.error(f"pptx_write 写入失败: {e}")
            return NodeResult(status=NodeStatus.FAILED, error=f"PPT 写入失败: {e}", summary="写入失败")

        logger.info(f"pptx_write 完成: {out_path} ({len(prs.slides)} slides)")
        return NodeResult(
            status=NodeStatus.SUCCESS,
            data={"file_path": str(out_path), "slide_count": len(prs.slides), "size": out_path.stat().st_size},
            summary=f"已生成 {out_path.name} ({len(prs.slides)} 页)",
        )
