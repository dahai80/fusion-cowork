from __future__ import annotations

import logging
import time
from typing import Any

from .registry import Skill, SkillRegistry

logger = logging.getLogger(__name__)


async def _cleanup_handler(args: str = "") -> Any:
    from ..engine.node import NodeConfig, NodeRegistry

    node = NodeRegistry.create(
        "desktop_clean",
        config=NodeConfig(
            params={
                "organize_by_type": True,
                "skip_hidden": True,
                "dry_run": False,
            }
        ),
    )
    if node is None:
        return {"error": "desktop_clean 节点不可用"}
    result = await node.execute({})
    return {"status": result.status.value, "summary": result.summary}


async def _classify_handler(args: str = "") -> Any:
    from ..engine.node import NodeConfig, NodeRegistry

    path = args.strip() or "~/Desktop"
    node = NodeRegistry.create(
        "ai_classify",
        config=NodeConfig(
            params={
                "classify_by_content": True,
            }
        ),
    )
    if node is None:
        return {"error": "ai_classify 节点不可用"}
    result = await node.execute({"path": path})
    return {"status": result.status.value, "summary": result.summary}


async def _screenshot_handler(args: str = "") -> Any:
    from ..engine.node import NodeConfig, NodeRegistry

    node = NodeRegistry.create(
        "screen_capture",
        config=NodeConfig(
            params={
                "save_to_clipboard": True,
            }
        ),
    )
    if node is None:
        return {"error": "screen_capture 节点不可用"}
    result = await node.execute({})
    return {"status": result.status.value, "summary": result.summary, "data": result.data}


async def _search_handler(args: str = "") -> Any:
    from ..engine.node import NodeConfig, NodeRegistry

    query = args.strip()
    if not query:
        return {"error": "请提供搜索关键词"}
    node = NodeRegistry.create(
        "web_search",
        config=NodeConfig(
            params={
                "query": query,
                "max_results": 5,
            }
        ),
    )
    if node is None:
        return {"error": "web_search 节点不可用"}
    result = await node.execute({"query": query})
    return {"status": result.status.value, "data": result.data}


async def _organize_handler(args: str = "") -> Any:
    from ..engine.node import NodeConfig, NodeRegistry

    node = NodeRegistry.create(
        "download_organizer",
        config=NodeConfig(
            params={
                "organize_by_type": True,
                "deduplicate": True,
                "dry_run": False,
            }
        ),
    )
    if node is None:
        return {"error": "download_organizer 节点不可用"}
    result = await node.execute({})
    return {"status": result.status.value, "summary": result.summary}


async def _diskclean_handler(args: str = "") -> Any:
    from ..engine.node import NodeConfig, NodeRegistry

    node = NodeRegistry.create(
        "disk_cleaner",
        config=NodeConfig(
            params={
                "clean_cache": True,
                "clean_temp": True,
                "clean_pycache": True,
                "clean_ds_store": True,
                "dry_run": False,
            }
        ),
    )
    if node is None:
        return {"error": "disk_cleaner 节点不可用"}
    result = await node.execute({})
    return {"status": result.status.value, "summary": result.summary}


async def _spreadsheet_recon_handler(args: str = "") -> Any:
    """issue #151: /spreadsheet-recon — 导出对账。args: "<budget.xlsx> <exports_dir>"。

    预算文件 sheet 前两列 = (科目, 预算额); exports_dir 下 *.xlsx/*.csv 各前两列 =
    (科目, 实际额, 同科目多文件累加)。产出对账工作簿 (差异 + 超阈值打标) +
    未对齐项备注, 全程本地确定性计算。
    """
    import csv as _csv
    from pathlib import Path

    from ..security import get_scoped_folder_manager

    parts = args.split()
    if len(parts) < 2:
        return {"error": "用法: /spreadsheet-recon <budget.xlsx> <exports_dir> [threshold_pct=10]"}
    budget_path = Path(parts[0]).expanduser()
    exports_dir = Path(parts[1]).expanduser()
    threshold_pct = float(parts[2]) if len(parts) > 2 else 10.0
    if not budget_path.is_file():
        return {"error": f"预算文件不存在: {budget_path}"}
    if not exports_dir.is_dir():
        return {"error": f"导出目录不存在: {exports_dir}"}
    scope = get_scoped_folder_manager()
    for p in (budget_path, exports_dir):
        if not scope.ensure_allowed(p):
            return {"error": f"沙箱拒绝越界读取: {p}"}

    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill
    except ImportError as e:
        return {"error": f"缺少依赖 openpyxl ({e}) — 安装: pip install -e '.[office]'"}

    def _num(v):
        try:
            return float(v)
        except (TypeError, ValueError):
            return None

    budget: dict = {}
    wb_budget = openpyxl.load_workbook(budget_path, data_only=True)
    ws_budget = wb_budget.active
    for row in ws_budget.iter_rows(min_row=2, values_only=True):
        if not row or row[0] in (None, ""):
            continue
        amount = _num(row[1]) if len(row) > 1 else None
        if amount is not None:
            budget[str(row[0]).strip()] = amount

    actual: dict = {}
    export_files = sorted(exports_dir.glob("*.xlsx")) + sorted(exports_dir.glob("*.csv"))
    for f in export_files:
        if f.suffix == ".xlsx":
            wb_f = openpyxl.load_workbook(f, data_only=True)
            rows = list(wb_f.active.iter_rows(min_row=2, values_only=True))
        else:
            with open(f, encoding="utf-8") as fh:
                rows = list(_csv.reader(fh))[1:]
        for row in rows:
            if not row or row[0] in (None, ""):
                continue
            amount = _num(row[1]) if len(row) > 1 else None
            if amount is not None:
                key = str(row[0]).strip()
                actual[key] = actual.get(key, 0.0) + amount

    matched, unreconciled_budget, unreconciled_actual = [], [], []
    for key, b in budget.items():
        if key in actual:
            a = actual[key]
            variance = a - b
            pct = (variance / abs(b) * 100) if b else (0.0 if variance == 0 else float("inf"))
            matched.append([key, b, a, round(variance, 2), "" if abs(pct) <= threshold_pct else "⚠️ 超阈值"])
        else:
            unreconciled_budget.append(key)
    for key in sorted(set(actual) - set(budget)):
        unreconciled_actual.append(key)

    out_path = exports_dir / f"reconciliation_{time.strftime('%Y%m%d_%H%M%S')}.xlsx"
    if not scope.ensure_allowed(out_path):
        return {"error": f"沙箱拒绝越界写入: {out_path}"}
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "对账结果"
    headers = ["科目", "预算", "实际", "差异", "标记"]
    ws.append(headers)
    flag_fill = PatternFill(start_color="FFFFC7CE", end_color="FFFFC7CE", fill_type="solid")
    for col in range(1, len(headers) + 1):
        c = ws.cell(row=1, column=col)
        c.font = Font(bold=True, color="FFFFFFFF")
        c.fill = PatternFill(start_color="FF305496", end_color="FF305496", fill_type="solid")
    for row in matched:
        ws.append(row)
        if row[4]:
            for col in range(1, len(headers) + 1):
                ws.cell(row=ws.max_row, column=col).fill = flag_fill
    wb.save(out_path)

    note = f"未对齐预算科目: {', '.join(unreconciled_budget) or '无'}; 无预算的实报科目: {', '.join(unreconciled_actual) or '无'}"
    flagged = sum(1 for r in matched if r[4])
    logger.info(f"spreadsheet-recon: {len(matched)} 项对账, {flagged} 项超 {threshold_pct}% 阈值, 输出 {out_path}")
    return {
        "status": "success",
        "summary": f"对账 {len(matched)} 项, {flagged} 项超阈值, 输出 {out_path.name}",
        "data": {
            "file_path": str(out_path),
            "matched": len(matched),
            "flagged": flagged,
            "unreconciled_budget": unreconciled_budget,
            "unreconciled_actual": unreconciled_actual,
            "note": note,
        },
    }


BUILTIN_SKILLS = [
    Skill(
        name="/cleanup",
        description="一键桌面清理 — 按类型自动规整桌面文件",
        handler=_cleanup_handler,
        category="desktop",
        aliases=["清理桌面", "桌面规整"],
    ),
    Skill(
        name="/classify",
        description="AI 文件分类 — 根据文件内容语义智能分类",
        handler=_classify_handler,
        category="ai",
        aliases=["文件分类", "智能分类"],
    ),
    Skill(
        name="/screenshot",
        description="截屏 — 截取当前屏幕并保存到剪贴板",
        handler=_screenshot_handler,
        category="desktop",
        aliases=["截屏", "屏幕截图"],
    ),
    Skill(
        name="/search",
        description="网页搜索 — 搜索互联网信息",
        handler=_search_handler,
        category="tool",
        aliases=["搜索", "网页搜索"],
    ),
    Skill(
        name="/organize",
        description="下载整理 — 自动归档下载文件夹",
        handler=_organize_handler,
        category="desktop",
        aliases=["下载整理", "下载归档"],
    ),
    Skill(
        name="/diskclean",
        description="磁盘清理 — 清理缓存、临时文件等垃圾",
        handler=_diskclean_handler,
        category="desktop",
        aliases=["磁盘清理", "清理磁盘"],
    ),
    Skill(
        name="/spreadsheet-recon",
        description="导出对账 — 预算 vs 导出目录差异对账, 产出对账工作簿 (issue #151)",
        handler=_spreadsheet_recon_handler,
        category="office",
        aliases=["导出对账", "对账"],
    ),
]


def register_builtin_skills(registry: SkillRegistry) -> None:
    for skill in BUILTIN_SKILLS:
        registry.register(skill)
    logger.info(f"注册 {len(BUILTIN_SKILLS)} 个内置技能")
