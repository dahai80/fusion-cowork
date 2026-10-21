"""issue #151 + #152: xlsx_write / pptx_write 交付物节点 + 对账技能 + 周报模板。

Contract:
- xlsx_write: 多 sheet、公式 cell ("=..." 打开时重算)、表头样式 (加粗+填充+冻结)、
  条件格式规则落盘、列宽; spec 缺失/无 sheets 显式失败。
- pptx_write: 标题/内容/图表/表格版式, period_compare 自动生成变动页 (>threshold%
  打标 ⚠️), 产物 python-pptx 可回读 (可编辑)。
- /spreadsheet-recon 技能: 预算 vs 导出目录对账, 差异超阈值打标 + 未对齐备注,
  产出对账工作簿。
- weekly_readout_pptx 模板: WorkflowEngine 端到端离线跑通 (本地 fixture spec)。
"""

from __future__ import annotations

import pytest

from fusion_cowork.engine.node import NodeConfig, NodeRegistry
from fusion_cowork.nodes.io.deliverables import PptxWriteNode as PptxCls
from fusion_cowork.nodes.io.deliverables import XlsxWriteNode as XlsxCls


@pytest.fixture(autouse=True)
def _register_nodes():
    saved = dict(NodeRegistry._registry)
    for cls in (XlsxCls, PptxCls):
        if cls.name not in NodeRegistry._registry:
            NodeRegistry._registry[cls.name] = cls
    yield
    NodeRegistry._registry.clear()
    NodeRegistry._registry.update(saved)


def _node(name: str, params: dict):
    node = NodeRegistry.create(name, config=NodeConfig(params=params))
    assert node is not None, f"节点 {name} 未注册"
    return node


class TestXlsxWrite:
    @pytest.mark.asyncio
    async def test_multi_sheet_formula_style_conditional(self, tmp_path):
        spec = {
            "sheets": [
                {
                    "name": "明细",
                    "headers": ["科目", "预算", "实际"],
                    "rows": [
                        ["营销", 100, 128],
                        ["研发", 200, 190],
                        ["差旅", 50, "=B3*2"],
                    ],
                    "column_widths": [20, 12, 12],
                    "conditional": [{"range": "C2:C4", "op": ">", "value": 150, "color": "FFFFC7CE"}],
                },
                {"name": "汇总", "rows": [["合计", "=SUM(B2:B4)", "=SUM(C2:C4)"]], "header_style": False},
            ],
            "active_sheet": "汇总",
        }
        node = _node("xlsx_write", {"output_path": str(tmp_path), "file_name": "wb_test", "overwrite": True})
        result = await node.execute({"spec": spec})
        assert result.status.value == "success", result.error
        out = tmp_path / "wb_test.xlsx"
        assert out.is_file()

        import openpyxl

        wb = openpyxl.load_workbook(out)
        assert wb.sheetnames == ["明细", "汇总"]
        ws = wb["明细"]
        assert ws.cell(row=1, column=1).value == "科目"
        assert ws.cell(row=1, column=1).font.bold, "表头应加粗"
        assert ws.freeze_panes == "A2", "表头应冻结"
        assert ws.cell(row=4, column=3).value == "=B3*2", "公式字符串应按公式写入"
        assert ws.column_dimensions["A"].width == 20
        cf_ranges = [str(rng) for rng in ws.conditional_formatting]
        assert any("C2:C4" in r for r in cf_ranges), f"条件格式应落盘: {cf_ranges}"
        assert (
            wb["汇总"].cell(row=1, column=1).font.bold is not True or wb["汇总"].cell(row=2, column=1).value == "合计"
        )

    @pytest.mark.asyncio
    async def test_missing_spec_fails(self, tmp_path):
        node = _node("xlsx_write", {"output_path": str(tmp_path)})
        result = await node.execute({})
        assert result.status.value == "failed"
        assert "spec" in (result.error or "")

    @pytest.mark.asyncio
    async def test_spec_json_string_accepted(self, tmp_path):
        import json

        node = _node(
            "xlsx_write",
            {
                "output_path": str(tmp_path),
                "file_name": "wb_str",
                "spec": json.dumps({"sheets": [{"name": "S", "rows": [[1, 2]]}]}),
            },
        )
        result = await node.execute({})
        assert result.status.value == "success"
        assert result.data["sheet_count"] == 1


class TestPptxWrite:
    @pytest.mark.asyncio
    async def test_full_deck_with_period_compare(self, tmp_path):
        spec = {
            "title": "周度读出",
            "subtitle": "本期 vs 上期",
            "slides": [
                {"layout": "content", "title": "要点", "bullets": ["第一条", "第二条"]},
                {
                    "layout": "chart",
                    "title": "对比",
                    "chart": {
                        "type": "column",
                        "categories": ["营收", "订单"],
                        "series": [{"name": "本期", "values": [120, 80]}],
                    },
                },
                {"layout": "table", "title": "明细", "headers": ["A", "B"], "rows": [[1, 2]]},
            ],
            "period_compare": {
                "metric": "指标",
                "categories": ["营收", "订单", "活跃"],
                "previous": [100, 90, 180],
                "current": [128, 84, 185],
                "threshold_pct": 10,
            },
        }
        node = _node("pptx_write", {"output_path": str(tmp_path), "file_name": "deck_test", "overwrite": True})
        result = await node.execute({"spec": spec})
        assert result.status.value == "success", result.error
        out = tmp_path / "deck_test.pptx"
        assert out.is_file()

        from pptx import Presentation

        prs = Presentation(out)
        # 标题页 + 3 slide + period_compare 2 页
        assert len(prs.slides) == 6
        slide_count = len(prs.slides)
        has_chart = any(getattr(shape, "has_chart", False) for shape in prs.slides[2].shapes)
        assert has_chart, "图表页应含真实图表对象 (可编辑)"
        # period_compare 变动表: 营收 +28% 打标, 订单 -6.7% 不打标
        text_all = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    text_all.append(shape.text_frame.text)
                if getattr(shape, "has_table", False):
                    for row in shape.table.rows:
                        text_all.append(" | ".join(c.text for c in row.cells))
        joined = "\n".join(text_all)
        assert "⚠️ 超阈值" in joined
        assert slide_count >= 5

    @pytest.mark.asyncio
    async def test_missing_spec_fails(self, tmp_path):
        node = _node("pptx_write", {"output_path": str(tmp_path)})
        result = await node.execute({})
        assert result.status.value == "failed"

    @pytest.mark.asyncio
    async def test_period_compare_zero_previous_no_crash(self, tmp_path):
        spec = {
            "period_compare": {
                "categories": ["新科目"],
                "previous": [0],
                "current": [10],
                "threshold_pct": 10,
            }
        }
        node = _node("pptx_write", {"output_path": str(tmp_path), "file_name": "deck_zero", "overwrite": True})
        result = await node.execute({"spec": spec})
        assert result.status.value == "success"


class TestSpreadsheetReconSkill:
    @pytest.mark.asyncio
    async def test_recon_end_to_end(self, tmp_path):
        import openpyxl

        budget = tmp_path / "budget.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["科目", "预算"])
        ws.append(["营销", 100])
        ws.append(["研发", 200])
        ws.append(["法务", 50])
        wb.save(budget)

        exports = tmp_path / "exports"
        exports.mkdir()
        (exports / "week1.csv").write_text("科目,实际\n营销,128\n研发,190\n", encoding="utf-8")
        (exports / "week2.csv").write_text("科目,实际\n营销,10\n", encoding="utf-8")
        wb2 = openpyxl.Workbook()
        ws2 = wb2.active
        ws2.append(["科目", "实际"])
        ws2.append(["行政", 30])
        wb2.save(exports / "misc.xlsx")

        from fusion_cowork.skills.builtin import _spreadsheet_recon_handler

        result = await _spreadsheet_recon_handler(f"{budget} {exports} 10")
        assert "error" not in result, result
        data = result["data"]
        assert data["matched"] == 2
        assert data["flagged"] == 1, "营销 138 vs 100 (+38%) 应超 10% 阈值"
        assert data["unreconciled_budget"] == ["法务"]
        assert data["unreconciled_actual"] == ["行政"]
        assert "法务" in data["note"] and "行政" in data["note"]

        out = openpyxl.load_workbook(data["file_path"])
        rows = list(out.active.iter_rows(values_only=True))
        assert rows[0] == ("科目", "预算", "实际", "差异", "标记")
        flagged = [r for r in rows[1:] if r[4]]
        assert len(flagged) == 1 and flagged[0][0] == "营销"

    @pytest.mark.asyncio
    async def test_missing_args(self):
        from fusion_cowork.skills.builtin import _spreadsheet_recon_handler

        result = await _spreadsheet_recon_handler("only-one-arg")
        assert "error" in result


class TestWeeklyReadoutTemplate:
    @pytest.mark.asyncio
    async def test_template_runs_offline(self, tmp_path):
        from fusion_cowork.engine.workflow import Workflow, WorkflowEngine
        from fusion_cowork.templates.template_manager import TemplateManager

        mgr = TemplateManager()
        tpl = mgr.get_template("weekly_readout_pptx")
        assert tpl is not None, "weekly_readout_pptx 模板应存在"
        wf = Workflow.from_dict(tpl["workflow"])
        engine = WorkflowEngine()
        result = await engine.execute(wf)
        assert result.status.value in ("success", "completed"), getattr(result, "error", None) or result.summary
        pptx_node = next(n for n in wf.nodes.values() if n.name == "pptx_write")
        assert pptx_node.config.params.get("spec"), "模板应内置本地 fixture spec"

    def test_registration_and_aliases(self):
        assert NodeRegistry.get("xlsx_write") is not None
        assert NodeRegistry.get("pptx_write") is not None
        from fusion_cowork import NODE_NAME_ALIASES

        assert NODE_NAME_ALIASES.get("写入Excel") == "xlsx_write"
        assert NODE_NAME_ALIASES.get("写入PPT") == "pptx_write"
